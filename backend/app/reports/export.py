"""Export hasil laporan ke CSV dan XLSX di memori (stateless, tanpa file di disk).

Teks yang diawali karakter formula (= + - @, tab, CR) tidak boleh dieksekusi spreadsheet:
di CSV diberi awalan apostrof (rekomendasi OWASP), di XLSX dipaksa bertipe teks.
"""

import csv
import io
from datetime import date, datetime
from decimal import Decimal
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font

from app.reports.catalog import ColumnDef, ColumnType
from app.schemas.reports import ReportResultColumn

FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def _is_formula_like(value: Any) -> bool:
    return isinstance(value, str) and value.startswith(FORMULA_PREFIXES)


def display_value(source: ColumnDef | None, value: Any) -> Any:
    """Nilai untuk manusia: label enum, Ya/Tidak, angka bulat tanpa desimal."""
    if value is None:
        return None
    if source is not None and source.type == ColumnType.ENUM:
        return source.enum_label(value)
    if isinstance(value, bool):
        return "Ya" if value else "Tidak"
    if isinstance(value, Decimal):
        return int(value) if value == value.to_integral_value() else float(value)
    if isinstance(value, datetime) and value.tzinfo is not None:
        return value.replace(tzinfo=None)  # Excel tidak mendukung zona waktu
    return value


def to_csv(
    columns: list[ReportResultColumn], sources: list[ColumnDef | None], rows: list[tuple[Any, ...]]
) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow([c.label for c in columns])
    for row in rows:
        values = []
        for source, raw in zip(sources, row, strict=True):
            value = display_value(source, raw)
            if isinstance(value, datetime | date):
                value = value.isoformat()
            if _is_formula_like(value):
                value = f"'{value}"
            values.append("" if value is None else value)
        writer.writerow(values)
    return buffer.getvalue().encode("utf-8-sig")  # BOM supaya Excel membaca UTF-8


def to_xlsx(
    title: str,
    columns: list[ReportResultColumn],
    sources: list[ColumnDef | None],
    rows: list[tuple[Any, ...]],
    definition: list[str],
) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Laporan"
    sheet.append([c.label for c in columns])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row_index, row in enumerate(rows, start=2):
        for col_index, (column, source, raw) in enumerate(
            zip(columns, sources, row, strict=True), start=1
        ):
            value = display_value(source, raw)
            cell = sheet.cell(row=row_index, column=col_index, value=value)
            if _is_formula_like(value):
                cell.data_type = "s"
            elif column.type == ColumnType.DATE and isinstance(value, date):
                cell.number_format = "yyyy-mm-dd"
            elif isinstance(value, datetime):
                cell.number_format = "yyyy-mm-dd hh:mm"
    sheet.freeze_panes = "A2"

    info = workbook.create_sheet("Definisi")
    info.append([title])
    info["A1"].font = Font(bold=True)
    for line in definition:
        info.append([line])

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
